"""
validate_brand.py — Auditor de cumplimiento de marca

Inspecciona un .pptx generado y reporta cualquier desviación de:
  - Tipografía (debe ser ForFuture Sans en todos los runs)
  - Paleta de colores (debe ser de colors.json)
  - Layouts (todos del catálogo)
  - Footer (presente y con formato correcto)

Uso:
  python validate_brand.py <deck.pptx> [--report report.md]
"""
from __future__ import annotations
import argparse
import json
import re
import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
except ImportError:
    print("ERROR: python-pptx is required.", file=sys.stderr)
    sys.exit(1)

SCRIPT_DIR = Path(__file__).resolve().parent
SKILL_DIR = SCRIPT_DIR.parent

with open(SKILL_DIR / "brand" / "colors.json") as f:
    COLORS = json.load(f)
with open(SKILL_DIR / "brand" / "typography.json") as f:
    TYPOGRAPHY = json.load(f)
with open(SKILL_DIR / "brand" / "layout_catalog.json") as f:
    LAYOUT_CATALOG = json.load(f)

ALLOWED_HEX = set()
for group in ("primary", "accents"):
    for key, color in COLORS[group].items():
        ALLOWED_HEX.add(color["hex"].lstrip("#").upper())
ALLOWED_HEX.add("000000")  # default black often appears
ALLOWED_HEX.add("FFFFFF")

ALLOWED_FONTS = {TYPOGRAPHY["family"]}
ALLOWED_LAYOUTS = {layout["name"] for layout in LAYOUT_CATALOG["layouts"]}

FOOTER_PATTERN = re.compile(r"MINSAIT\s*•\s*.+?\s*•\s*\d{2}/\d{2}/\d{4}")


def validate(pptx_path: str) -> dict:
    """Validate a .pptx and return a report dict."""
    prs = Presentation(pptx_path)
    report = {
        "file": pptx_path,
        "total_slides": len(prs.slides),
        "issues": {
            "fonts": [],
            "colors": [],
            "layouts": [],
            "footer": [],
        },
        "summary": {},
    }

    for i, slide in enumerate(prs.slides, start=1):
        layout_name = slide.slide_layout.name or ""
        if layout_name and layout_name not in ALLOWED_LAYOUTS:
            report["issues"]["layouts"].append(
                f"Slide {i}: usa layout '{layout_name}' fuera del catálogo oficial."
            )

        slide_text = ""
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    slide_text += run.text + " "

                    # Check font
                    font_name = run.font.name
                    if font_name and font_name not in ALLOWED_FONTS:
                        report["issues"]["fonts"].append(
                            f"Slide {i}: fuente no oficial '{font_name}' en texto: '{run.text[:40]}'"
                        )

                    # Check color (only if explicitly set, not theme-derived)
                    try:
                        if run.font.color and run.font.color.type and run.font.color.rgb:
                            hex_color = str(run.font.color.rgb).upper()
                            if hex_color not in ALLOWED_HEX:
                                report["issues"]["colors"].append(
                                    f"Slide {i}: color no oficial #{hex_color} en texto: '{run.text[:40]}'"
                                )
                    except Exception:
                        pass

        # Footer check (skip first and last slide, which are cover and closing)
        if 1 < i < report["total_slides"]:
            if not FOOTER_PATTERN.search(slide_text):
                report["issues"]["footer"].append(
                    f"Slide {i}: footer no detectado o con formato no estándar."
                )

    report["summary"] = {
        "fonts_issues": len(report["issues"]["fonts"]),
        "color_issues": len(report["issues"]["colors"]),
        "layout_issues": len(report["issues"]["layouts"]),
        "footer_issues": len(report["issues"]["footer"]),
        "total_issues": sum(len(v) for v in report["issues"].values()),
    }
    return report


def render_report(report: dict) -> str:
    """Render a markdown report from a validation result."""
    lines = [f"# Reporte de Validación de Marca · {Path(report['file']).name}", ""]
    lines.append(f"**Total slides:** {report['total_slides']}")
    lines.append(f"**Total issues:** {report['summary']['total_issues']}")
    lines.append("")

    if report["summary"]["total_issues"] == 0:
        lines.append("## ✅ El deck cumple las reglas de marca")
        return "\n".join(lines)

    lines.append("## ⚠️ Issues detectados")
    for category, issues in report["issues"].items():
        if not issues:
            continue
        lines.append(f"### {category.title()} ({len(issues)})")
        for issue in issues:
            lines.append(f"- {issue}")
        lines.append("")
    return "\n".join(lines)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("pptx", help="Path to the .pptx to validate")
    parser.add_argument("--report", help="Output markdown report path", default=None)
    args = parser.parse_args()

    report = validate(args.pptx)
    md = render_report(report)
    print(md)
    if args.report:
        Path(args.report).write_text(md)
    sys.exit(0 if report["summary"]["total_issues"] == 0 else 1)


if __name__ == "__main__":
    main()
